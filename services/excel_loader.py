"""
Chargement et validation du fichier Excel d'entrée.
Format strict : 2 feuilles obligatoires (Postes, Parametres).
"""
from typing import List, Tuple
from io import BytesIO
import openpyxl
from .domain import Poste, Parametres


class ExcelStructureError(Exception):
    """Erreur de structure du fichier Excel."""
    pass


REQUIRED_POSTES_COLS = ["numero", "nom_operation", "machine", "temps_min_par_piece"]
REQUIRED_PARAM_KEYS = [
    "demande_client_jour", "temps_travail_min_jour",
    "taux_rebut", "taux_arret_machine", "taille_lot",
]


def charger_excel(file_bytes: bytes) -> Tuple[List[Poste], Parametres]:
    """
    Charge un Excel et retourne (postes, parametres).
    Lève ExcelStructureError si le format est invalide.
    """
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ExcelStructureError(f"Fichier Excel illisible : {e}")

    sheet_names = [s.lower() for s in wb.sheetnames]
    if "postes" not in sheet_names:
        raise ExcelStructureError("Feuille 'Postes' manquante.")
    if "parametres" not in sheet_names:
        raise ExcelStructureError("Feuille 'Parametres' manquante.")

    # Récupérer les feuilles avec correspondance insensible à la casse
    ws_postes = wb[wb.sheetnames[sheet_names.index("postes")]]
    ws_params = wb[wb.sheetnames[sheet_names.index("parametres")]]

    postes = _charger_postes(ws_postes)
    parametres = _charger_parametres(ws_params)

    return postes, parametres


def _charger_postes(ws) -> List[Poste]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ExcelStructureError("Feuille 'Postes' vide.")

    headers = [str(c).strip().lower() if c else "" for c in rows[0]]
    for col in REQUIRED_POSTES_COLS:
        if col not in headers:
            raise ExcelStructureError(f"Colonne '{col}' manquante dans feuille 'Postes'.")

    idx = {col: headers.index(col) for col in REQUIRED_POSTES_COLS}
    idx_nb = headers.index("nb_machines") if "nb_machines" in headers else None

    postes: List[Poste] = []
    for r, row in enumerate(rows[1:], start=2):
        if all(c is None or c == "" for c in row):
            continue
        try:
            poste = Poste(
                numero=int(row[idx["numero"]]),
                nom_operation=str(row[idx["nom_operation"]]),
                machine=str(row[idx["machine"]]),
                temps_min_par_piece=float(row[idx["temps_min_par_piece"]]),
                nb_machines=int(row[idx_nb]) if idx_nb is not None and row[idx_nb] else 1,
            )
        except (ValueError, TypeError) as e:
            raise ExcelStructureError(f"Ligne {r} de 'Postes' invalide : {e}")
        if poste.temps_min_par_piece <= 0:
            raise ExcelStructureError(f"Ligne {r} : temps doit être > 0.")
        if poste.nb_machines < 1:
            raise ExcelStructureError(f"Ligne {r} : nb_machines doit être >= 1.")
        postes.append(poste)

    if not postes:
        raise ExcelStructureError("Aucun poste valide trouvé.")

    return postes


def _charger_parametres(ws) -> Parametres:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ExcelStructureError("Feuille 'Parametres' vide.")

    headers = [str(c).strip().lower() if c else "" for c in rows[0]]
    if "cle" not in headers or "valeur" not in headers:
        raise ExcelStructureError("Feuille 'Parametres' doit avoir colonnes 'cle' et 'valeur'.")

    i_cle = headers.index("cle")
    i_val = headers.index("valeur")

    data = {}
    for row in rows[1:]:
        if not row or row[i_cle] is None:
            continue
        cle = str(row[i_cle]).strip().lower()
        try:
            data[cle] = float(row[i_val])
        except (ValueError, TypeError):
            raise ExcelStructureError(f"Valeur invalide pour paramètre '{cle}'.")

    for key in REQUIRED_PARAM_KEYS:
        if key not in data:
            raise ExcelStructureError(f"Paramètre obligatoire '{key}' manquant.")

    # Validation sémantique
    if not (0 <= data["taux_rebut"] <= 1):
        raise ExcelStructureError("taux_rebut doit être entre 0 et 1.")
    if not (0 <= data["taux_arret_machine"] <= 1):
        raise ExcelStructureError("taux_arret_machine doit être entre 0 et 1.")

    return Parametres(
        demande_client_jour=data["demande_client_jour"],
        temps_travail_min_jour=data["temps_travail_min_jour"],
        taux_rebut=data["taux_rebut"],
        taux_arret_machine=data["taux_arret_machine"],
        taille_lot=int(data["taille_lot"]),
        prix_vente_unitaire=data.get("prix_vente_unitaire", 85.0),
        cout_variable_unitaire=data.get("cout_variable_unitaire", 48.0),
        jours_ouvres_an=int(data.get("jours_ouvres_an", 220)),
    )


def generer_template_excel() -> bytes:
    """Génère un template Excel pré-rempli avec les valeurs du PDF."""
    wb = openpyxl.Workbook()

    # Feuille Postes
    ws1 = wb.active
    ws1.title = "Postes"
    ws1.append(["numero", "nom_operation", "machine", "temps_min_par_piece", "nb_machines"])
    donnees_pdf = [
        (1, "Découpe brut",       "Scie",       2.0, 1),
        (2, "Tournage extérieur", "Tour",       6.0, 1),
        (3, "Perçage axial",      "Perceuse",   4.0, 1),
        (4, "Fraisage rainure",   "Fraiseuse",  5.0, 1),
        (5, "Finition",           "Rectifieuse", 3.0, 1),
        (6, "Contrôle qualité",   "Manuel",     2.0, 1),
    ]
    for row in donnees_pdf:
        ws1.append(row)

    # Feuille Parametres
    ws2 = wb.create_sheet("Parametres")
    ws2.append(["cle", "valeur"])
    params = [
        ("demande_client_jour",   80),
        ("temps_travail_min_jour", 480),
        ("taux_rebut",            0.05),
        ("taux_arret_machine",    0.10),
        ("taille_lot",            20),
        ("prix_vente_unitaire",   85),
        ("cout_variable_unitaire", 48),
        ("jours_ouvres_an",       220),
    ]
    for row in params:
        ws2.append(row)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
